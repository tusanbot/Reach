import csv
import io
from pathlib import Path

from aiogram import F, Router
from aiogram.fsm.context import FSMContext
from aiogram.types import InlineKeyboardButton, InlineKeyboardMarkup

from app.handlers.builder import BuilderState

router = Router()

TITLES = {"table": "📊 جدول", "ranking": "🏆 رتبه‌بندی", "stats": "📈 آمار", "custom": "📝 پیام سفارشی"}

def _decode(raw):
    for enc in ("utf-8-sig", "utf-8", "cp1256", "windows-1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            pass
    raise ValueError("encoding فایل قابل تشخیص نیست.")

def _clean(rows):
    rows = [[("" if x is None else str(x)).strip() for x in row] for row in rows]
    rows = [row for row in rows if any(row)]
    if len(rows) < 2:
        raise ValueError("فایل باید حداقل یک ردیف عنوان و یک ردیف داده داشته باشد.")
    width = len(rows[0])
    if width < 2:
        raise ValueError("جدول باید حداقل دو ستون داشته باشد.")
    normalized = []
    for row in rows:
        if len(row) > width:
            raise ValueError("تعداد ستون‌های فایل یکسان نیست.")
        normalized.append(row + [""] * (width - len(row)))
    headers = [x or f"ستون {i}" for i, x in enumerate(normalized[0], 1)]
    return {"headers": headers, "rows": normalized[1:]}

def _parse_text(text):
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return _clean(list(csv.reader(io.StringIO(text), dialect)))

def _parse_xlsx(raw):
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    result = {}
    for name in book.sheetnames:
        rows = [list(row) for row in book[name].iter_rows(values_only=True)]
        try:
            result[name] = _clean(rows)
        except ValueError:
            continue
    book.close()
    if not result:
        raise ValueError("هیچ Sheet دارای جدول معتبر پیدا نشد.")
    return result

def _menu(names):
    rows = [[InlineKeyboardButton(text=f"📄 {name}", callback_data=f"sheet:{i}")] for i, name in enumerate(names)]
    rows.append([InlineKeyboardButton(text="❌ لغو", callback_data="file:cancel")])
    return InlineKeyboardMarkup(inline_keyboard=rows)

async def _apply(state, table):
    data = await state.get_data()
    kind = data.get("kind", "table")
    config = data.get("template_config") or {}
    await state.update_data(
        headers=table["headers"], rows=table["rows"],
        title=config.get("title", TITLES.get(kind, "📊 جدول")),
        subtitle=config.get("subtitle", ""),
        footer=config.get("footer", ""),
        style=config.get("style", "classic"),
        show_index=bool(config.get("show_index", kind == "ranking")),
        align=config.get("align", "left"),
        importing=False
    )
    await state.set_state(BuilderState.preview)

@router.callback_query(F.data == "file:upload")
async def ask_file(callback, state: FSMContext):
    await callback.answer()
    await state.update_data(importing=True)
    await state.set_state(BuilderState.waiting_data)
    await callback.message.answer("📎 فایل CSV، TXT یا XLSX را ارسال کن. ردیف اول عنوان ستون‌هاست.")

@router.message(BuilderState.waiting_data, F.document)
async def receive_file(message, state: FSMContext):
    data = await state.get_data()
    if not data.get("importing"):
        return
    document = message.document
    suffix = Path(document.file_name or "").suffix.lower()
    if suffix not in {".csv", ".txt", ".xlsx"}:
        await message.answer("⚠️ فقط CSV، TXT و XLSX پشتیبانی می‌شوند.")
        return
    try:
        raw = io.BytesIO()
        await message.bot.download(document, destination=raw)
        payload = raw.getvalue()
        if suffix == ".xlsx":
            sheets = _parse_xlsx(payload)
            if len(sheets) > 1:
                await state.update_data(import_sheets=sheets)
                await message.answer("📚 چند Sheet معتبر پیدا شد؛ یکی را انتخاب کن:", reply_markup=_menu(list(sheets)))
                return
            table = next(iter(sheets.values()))
        else:
            table = _parse_text(_decode(payload))
        await _apply(state, table)
        from app.services.renderer import render_message
        d = await state.get_data()
        preview = render_message(d["headers"], d["rows"], title=d["title"], style=d["style"], show_index=d["show_index"], align=d["align"])
        from app.keyboards import preview_menu
        await message.answer(preview, reply_markup=preview_menu())
    except Exception as exc:
        await message.answer(f"⚠️ ورود فایل انجام نشد: {exc}")

@router.callback_query(F.data.startswith("sheet:"))
async def select_sheet(callback, state: FSMContext):
    data = await state.get_data()
    sheets = data.get("import_sheets", {})
    try:
        index = int(callback.data.split(":", 1)[1])
        table = list(sheets.values())[index]
    except (ValueError, IndexError):
        await callback.answer("Sheet معتبر نیست", show_alert=True)
        return
    await _apply(state, table)
    await state.update_data(import_sheets=None)
    from app.services.renderer import render_message
    from app.keyboards import preview_menu
    d = await state.get_data()
    preview = render_message(d["headers"], d["rows"], title=d["title"], style=d["style"], show_index=d["show_index"], align=d["align"])
    await callback.answer("Sheet انتخاب شد")
    await callback.message.edit_text(preview, reply_markup=preview_menu())

@router.callback_query(F.data == "file:cancel")
async def cancel_file(callback, state: FSMContext):
    await state.update_data(importing=False, import_sheets=None)
    await callback.answer("لغو شد")
    from app.keyboards import main_menu
    await callback.message.edit_text("❌ ورود فایل لغو شد.", reply_markup=main_menu())
